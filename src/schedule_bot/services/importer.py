import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from schedule_bot.repositories import GroupRepository, LessonRepository
from schedule_bot.structure import normalize_faculty, normalize_specialty

ALIASES = {
    "group": {"group", "группа"},
    "faculty": {"faculty", "факультет", "институт"},
    "course": {"course", "курс"},
    "specialty": {"specialty", "направление", "специальность", "профиль"},
    "date": {"date", "дата"},
    "start_time": {"start_time", "начало", "время начала"},
    "end_time": {"end_time", "конец", "время конца"},
    "subject": {"subject", "предмет", "дисциплина"},
    "teacher": {"teacher", "преподаватель"},
    "room": {"room", "аудитория", "кабинет"},
    "lesson_type": {"type", "lesson_type", "тип", "вид занятия"},
}
REQUIRED = {"faculty", "course", "group", "date", "start_time", "subject"}
MAX_ROWS = 20_000


class ImportValidationError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class LessonRow:
    group: str
    faculty: str
    specialty: str | None
    course: int
    lesson_date: date
    start_time: time
    end_time: time | None
    subject: str
    teacher: str | None
    room: str | None
    lesson_type: str | None


@dataclass(frozen=True, slots=True)
class ImportResult:
    total: int
    created: int
    updated: int
    groups: int


def parse_schedule(filename: str, content: bytes) -> list[LessonRow]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        rows = _xlsx_rows(content)
    elif suffix == ".csv":
        rows = _csv_rows(content)
    else:
        raise ImportValidationError("Поддерживаются только файлы .xlsx и .csv")

    iterator = iter(rows)
    try:
        raw_headers = next(iterator)
    except StopIteration as error:
        raise ImportValidationError("Файл пуст") from error
    headers = _map_headers(raw_headers)

    parsed: list[LessonRow] = []
    errors: list[str] = []
    for number, values in enumerate(iterator, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        if len(parsed) >= MAX_ROWS:
            raise ImportValidationError(f"В одном файле допускается не более {MAX_ROWS} строк")
        raw = {key: values[index] if index < len(values) else None for key, index in headers.items()}
        try:
            parsed.append(_parse_row(raw))
        except (TypeError, ValueError) as error:
            errors.append(f"строка {number}: {error}")
            if len(errors) == 10:
                break

    if errors:
        raise ImportValidationError("Ошибки в файле:\n" + "\n".join(errors))
    if not parsed:
        raise ImportValidationError("В файле нет занятий")
    return parsed


def submission_identity(rows: list[LessonRow]) -> tuple[str, str, str | None, int]:
    identities = {
        (row.group, row.faculty, row.specialty, row.course)
        for row in rows
    }
    if len(identities) != 1:
        raise ImportValidationError(
            "В пользовательской заявке должно быть расписание только одной группы "
            "с одинаковыми факультетом, направлением и курсом."
        )
    return next(iter(identities))


async def save_schedule(session: AsyncSession, rows: list[LessonRow]) -> ImportResult:
    groups = GroupRepository(session)
    lessons = LessonRepository(session)
    cache = {}
    created = 0
    for row in rows:
        group = cache.get(row.group)
        if group is None:
            group = await groups.get_or_create(
                row.group, row.faculty, row.specialty, row.course
            )
            cache[row.group] = group
        is_created = await lessons.upsert(
            group_id=group.id,
            lesson_date=row.lesson_date,
            start_time=row.start_time,
            end_time=row.end_time,
            subject=row.subject,
            teacher=row.teacher,
            room=row.room,
            lesson_type=row.lesson_type,
        )
        created += int(is_created)
    await session.commit()
    return ImportResult(
        total=len(rows), created=created, updated=len(rows) - created, groups=len(cache)
    )


def _xlsx_rows(content: bytes) -> Iterable[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        return ([cell for cell in row] for row in sheet.iter_rows(values_only=True))
    except Exception as error:
        raise ImportValidationError("Не удалось прочитать XLSX-файл") from error


def _csv_rows(content: bytes) -> Iterable[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return csv.reader(io.StringIO(text), dialect)


def _map_headers(raw_headers: list[Any]) -> dict[str, int]:
    normalized = {str(value).strip().lower(): index for index, value in enumerate(raw_headers)}
    mapped = {
        canonical: normalized[alias]
        for canonical, aliases in ALIASES.items()
        for alias in aliases
        if alias in normalized
    }
    missing = REQUIRED - mapped.keys()
    if missing:
        raise ImportValidationError("Не найдены обязательные колонки: " + ", ".join(sorted(missing)))
    return mapped


def _parse_row(raw: dict[str, Any]) -> LessonRow:
    group = _required_text(raw.get("group"), "группа")
    subject = _required_text(raw.get("subject"), "предмет")
    faculty = normalize_faculty(raw.get("faculty"))
    specialty = normalize_specialty(raw.get("specialty"), faculty)
    course = _course(raw.get("course"))
    if course is None:
        raise ValueError("не заполнено поле «курс»")
    return LessonRow(
        group=group,
        faculty=faculty,
        specialty=specialty,
        course=course,
        lesson_date=_date(raw.get("date")),
        start_time=_time(raw.get("start_time"), "время начала"),
        end_time=_optional_time(raw.get("end_time")),
        subject=subject,
        teacher=_optional_text(raw.get("teacher")),
        room=_optional_text(raw.get("room")),
        lesson_type=_optional_text(raw.get("lesson_type")),
    )


def _required_text(value: Any, name: str) -> str:
    result = str(value).strip() if value is not None else ""
    if not result:
        raise ValueError(f"не заполнено поле «{name}»")
    return result


def _optional_text(value: Any) -> str | None:
    result = str(value).strip() if value is not None else ""
    return result or None


def _course(value: Any) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip().lower().replace("курс", "").strip()
    if not text.isdigit():
        raise ValueError("неверный курс (нужно число, например 2)")
    course = int(text)
    if not 1 <= course <= 9:
        raise ValueError("курс должен быть числом от 1 до 9")
    return course


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("неверная дата (нужно ДД.ММ.ГГГГ или ГГГГ-ММ-ДД)")


def _time(value: Any, name: str) -> time:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = str(value).strip()
    for pattern in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            continue
    raise ValueError(f"неверное поле «{name}» (нужно ЧЧ:ММ)")


def _optional_time(value: Any) -> time | None:
    if value in (None, ""):
        return None
    return _time(value, "время конца")
